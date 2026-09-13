r"""xlsx → records 迁移脚本（架构 §8.1 逐列映射 + §8.2 基线口径）。

用法（在项目根）：
    server\.venv\Scripts\python.exe scripts\import_xlsx.py

要点：
- 只读 assets\ 副本，绝不改源文件
- 只导"今天之前"且有真实数据的行（跳过预填未来空行）
- sheet+行号记入 imports 台账，幂等可重跑
- I/N 列为自由文本，按规则推断枚举档位（P1 验收口径："一大杯可乐"→全糖）
- 体重趋势 sheet 不导公式值（#DIV/0!），基线参数写入 users.settings_json
- 使用说明 sheet → kb_chunks（口径约定）
"""
import sys
from datetime import date, datetime, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SERVER = ROOT / "server"
sys.path.insert(0, str(SERVER))

import openpyxl  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app import schemas  # noqa: E402
from app.database import SessionLocal, make_engine  # noqa: E402
from app.models import ImportLedger, KbChunk, KbDoc, Record, User  # noqa: E402

XLSX = ROOT / "assets" / "august_fatloss" / "每日记录.xlsx"
SOURCE_FILE = "每日记录.xlsx"

# §8.2 基线口径（写入 users.settings_json）
BASELINE = {
    "initial_weight_kg": 108.55,
    "initial_weight_date": "2026-08-30",
    "initial_waist_cm": 110,
    "phases": {
        "sprint": {"name": "冲刺期", "start": "2026-08-31", "end": "2026-09-20"},
        "steady": {"name": "稳态期", "start": "2026-09-21"},
    },
    "targets": {
        "sprint_end": {"weight_kg_range": [104, 105], "waist_cm_range": [107, 108]},
        "long_term": {"weight_kg_range": [93, 95], "waist_cm_range": [102, 104]},
    },
}


def is_bad(v) -> bool:
    """空值或公式错误值。"""
    if v is None or v == "":
        return True
    return isinstance(v, str) and v.startswith("#")


def infer_sweet_level(text: str) -> str | None:
    """自由文本 → 奶茶档位枚举（取出现的最高档；"可乐"视为全糖）。"""
    if not text:
        return None
    if "全糖" in text or "可乐" in text:
        return "全糖"
    if "半糖" in text:
        return "半糖"
    if "三分糖" in text or "3分糖" in text:
        return "三分糖"
    if "无糖" in text:
        return "无糖"
    if "0杯" in text or text.strip() == "0":
        return "0杯"
    return None


def infer_night_hunger(text: str) -> str | None:
    """自由文本 → 夜饿处理枚举；有实际内容（吃了东西）= 破戒。"""
    if not text:
        return None
    if "没饿" in text:
        return "没饿"
    if "预案" in text or "加餐" in text:
        return "加餐预案"
    return "破戒"


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def as_bedtime(v) -> str | None:
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, str) and v.strip():
        return v.strip()
    return None


def build_records(row: int, d: date, cells: dict) -> list[dict]:
    """一行 xlsx → 若干待插入记录（先过 Pydantic 闸再入库）。"""
    out = []

    def add(rtype: str, slot: str | None, fields: dict, raw: str) -> None:
        try:  # 复用后端闸：范围/枚举/槽位校验
            model = schemas.FIELDS_MODEL[rtype](**fields)
            fields = model.model_dump(exclude_none=True)
        except Exception as e:  # noqa: BLE001 记台账继续
            print(f"  [跳过] 第{row}行 {rtype}: 校验失败 {e}")
            return
        out.append(
            {"type": rtype, "date": d, "slot": slot, "fields": fields, "raw_text": raw}
        )

    if not is_bad(cells.get("C")):
        add("weight", "am", {"kg": float(cells["C"])}, str(cells["C"]))
    if not is_bad(cells.get("D")):
        add("weight", "evening", {"kg": float(cells["D"])}, str(cells["D"]))
    for col, slot in zip("EFGH", ["breakfast", "lunch", "dinner", "snack"]):
        v = cells.get(col)
        if not is_bad(v):
            add("meal", slot, {"desc": str(v).strip()}, str(v).strip())
    v = cells.get("I")
    if not is_bad(v):
        text = str(v).strip()
        level = infer_sweet_level(text)
        if level:
            fields = {"level": level}
            if text != level:
                fields["desc"] = text
            add("sweet_drink", None, fields, text)
        else:
            print(f"  [跳过] 第{row}行 sweet_drink: 无法推断档位 {text!r}")
    v = cells.get("J")
    if not is_bad(v):
        add("exercise", None, {"kind": "swim", "duration_min": int(v)}, str(v))
    v = cells.get("K")
    if not is_bad(v):
        text = str(v).strip()
        if "完成" in text:
            add("exercise", None, {"kind": "strength", "done": True}, text)
        elif "未做" in text or "没做" in text:
            add("exercise", None, {"kind": "strength", "done": False}, text)
        else:
            print(f"  [跳过] 第{row}行 strength: 无法解析 {text!r}")
    v = cells.get("L")
    if not is_bad(v):
        add("step", None, {"steps": int(v)}, str(v))
    v = cells.get("M")
    if not is_bad(v):
        bt = as_bedtime(v)
        if bt:
            add("sleep", None, {"bedtime": bt}, bt)
    v = cells.get("N")
    if not is_bad(v):
        text = str(v).strip()
        level = infer_night_hunger(text)
        if level:
            add("night_hunger", None, {"level": level}, text)
        else:
            print(f"  [跳过] 第{row}行 night_hunger: 无法解析 {text!r}")
    v = cells.get("O")
    if not is_bad(v):
        add("note", None, {"text": str(v).strip()}, str(v).strip())
    return out


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    today = date.today()
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["每日记录"]

    db = SessionLocal()
    inserted = skipped_ledger = 0
    try:
        for r in range(3, ws.max_row + 1):
            d = as_date(ws.cell(row=r, column=1).value)
            if d is None or d >= today:  # 只导今天之前（§8.2）
                continue
            cells = {c: ws.cell(row=r, column=i).value for i, c in enumerate("ABCDEFGHIJKLMNO", start=1)}
            if all(is_bad(v) for k, v in cells.items() if k not in "AB"):
                continue  # 预填空行
            # 幂等：台账已有该行则跳过
            if db.scalars(
                select(ImportLedger).where(
                    ImportLedger.source_file == SOURCE_FILE,
                    ImportLedger.sheet == "每日记录",
                    ImportLedger.row_no == r,
                )
            ).first():
                skipped_ledger += 1
                continue
            recs = build_records(r, d, cells)
            for item in recs:
                db.add(
                    Record(
                        type=item["type"],
                        date=item["date"],
                        slot=item["slot"],
                        fields_json=item["fields"],
                        raw_text=item["raw_text"],
                        source="import",
                        confirmed=True,
                    )
                )
                inserted += 1
            db.add(
                ImportLedger(
                    source_file=SOURCE_FILE, sheet="每日记录", row_no=r,
                    note=f"{d} 导入 {len(recs)} 条",
                )
            )
            print(f"第{r}行 {d}: +{len(recs)} 条")
        db.commit()

        # 使用说明 sheet → kb_chunks（口径约定）
        n_kb = import_usage_sheet(db, wb)
        db.commit()

        # 基线参数 → users.settings_json
        write_baseline(db)
        db.commit()
    finally:
        db.close()

    print(f"\n完成：新入库 {inserted} 条 records；台账跳过 {skipped_ledger} 行；"
          f"使用说明 KB {n_kb} 块；基线参数已写入 settings。")


def import_usage_sheet(db, wb) -> int:
    """xlsx"使用说明"sheet → kb_chunks（§8.1：口径约定作 KB 补充）。"""
    ws = wb["使用说明"]
    lines = [
        str(ws.cell(row=r, column=1).value).strip()
        for r in range(1, ws.max_row + 1)
        if ws.cell(row=r, column=1).value not in (None, "")
    ]
    if not lines:
        return 0
    doc = db.scalars(
        select(KbDoc).where(
            KbDoc.title == "每日记录.xlsx 使用说明",
            KbDoc.source_plan == "august_fatloss",
        )
    ).first()
    if doc is None:
        doc = KbDoc(title="每日记录.xlsx 使用说明", source_plan="august_fatloss", version="2026-08-30")
        db.add(doc)
        db.flush()
    # 覆盖旧版（§6.4 版本化）：重跑时先删旧块再加新块
    for old in db.scalars(select(KbChunk).where(KbChunk.doc_id == doc.id)).all():
        db.delete(old)
    db.flush()
    db.add(
        KbChunk(
            doc_id=doc.id, title=doc.title, heading="使用说明（记录口径）",
            source_plan="august_fatloss", section_no="",
            content="\n".join(lines), tags=["监测"],
        )
    )
    return 1


def write_baseline(db) -> None:
    """体重趋势 sheet 的基线口径 → users.settings_json（不导公式值）。"""
    user = db.scalars(select(User).where(User.username == "me")).first()
    if user is None:
        user = User(username="me", settings_json={})
        db.add(user)
    settings = dict(user.settings_json or {})
    settings["baseline"] = BASELINE
    user.settings_json = settings


if __name__ == "__main__":
    main()
